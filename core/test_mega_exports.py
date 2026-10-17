from datetime import date, datetime
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core.models import Mahalla, UzchessMahallaStat, UzchessStatSnapshot, Yosh


class MegaExcelExportTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.mahalla = Mahalla.objects.create(name="Excel Mahalla")
        cls.user = get_user_model().objects.create_user(
            username="mega-export-user",
            password="test-password",
            role="RAHBAR",
        )
        for index in range(10):
            Yosh.objects.create(
                fullname=f"Excel Youth {index}",
                mahalla=cls.mahalla,
                birth_date=date(2000, 1, 1),
                jshshir=f"9988776655{index:04d}",
                address="Test address",
            )

        cls.first_snapshot = UzchessStatSnapshot.objects.create(
            snapshot_date=date(2026, 1, 1),
            source_url="https://example.test/old",
            raw_payload={},
            fetched_at=timezone.make_aware(datetime(2026, 1, 1, 9, 0)),
        )
        cls.last_snapshot = UzchessStatSnapshot.objects.create(
            snapshot_date=date(2026, 2, 1),
            source_url="https://example.test/new",
            raw_payload={},
            fetched_at=timezone.make_aware(datetime(2026, 2, 1, 9, 0)),
        )
        UzchessMahallaStat.objects.create(
            snapshot=cls.first_snapshot,
            mahalla=cls.mahalla,
            area_name=cls.mahalla.name,
            metrics={"profiles": 2, "submissions": 5, "games_count": 7, "certificates": 1},
        )
        UzchessMahallaStat.objects.create(
            snapshot=cls.last_snapshot,
            mahalla=cls.mahalla,
            area_name=cls.mahalla.name,
            metrics={"profiles": 8, "submissions": 9, "games_count": 12, "certificates": 3},
        )

    def setUp(self):
        self.client.force_login(self.user)

    def test_stats_page_renders_excel_configuration_dialog(self):
        response = self.client.get(reverse("mega_uzchess"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Excel yuklab olish")
        self.assertContains(response, "Jadvalga kiritiladigan snapshotlar")
        self.assertContains(response, 'name="snapshot_ids"')
        self.assertContains(response, "ikkita snapshot farqini ko")
        self.assertContains(response, "Foiz chegaralarini qo'lda belgilash")

    def test_exports_selected_snapshots_with_totals_deltas_and_row_color(self):
        # Default: 14-30 yoshlar at start, snapshots have only profiles, plus deltas and jami foiz
        response = self.client.post(
            reverse("mega_excel_export", kwargs={"project_key": "uzchess"}),
            {
                "snapshot_ids": [str(self.first_snapshot.pk), str(self.last_snapshot.pk)],
                "compare_latest": "on",
                "compare_first_last": "on",
                "color_mode": "traffic",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "attachment; filename=\"uzchess-statistika.xlsx\"",
            response["Content-Disposition"],
        )
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        worksheet = workbook.active
        self.assertEqual(worksheet["A5"].value, self.mahalla.name)
        self.assertEqual(worksheet["A6"].value, "Tuman bo'yicha jami")

        # Column B is 14-30 yoshlar placed once at the beginning
        self.assertEqual(worksheet["B3"].value, "14-30 yoshlar")
        self.assertEqual(worksheet["B5"].value, 10)
        self.assertEqual(worksheet["B6"].value, 10)

        # Snapshot columns by default only contain Profiles
        self.assertEqual(worksheet["C4"].value, "Profiles")
        self.assertEqual(worksheet["C5"].value, 2)
        self.assertEqual(worksheet["D4"].value, "Profiles")
        self.assertEqual(worksheet["D5"].value, 8)

        # Delta columns for Profiles
        self.assertEqual(worksheet["E4"].value, "Farq (Profiles)")
        self.assertEqual(worksheet["E5"].value, 6)
        self.assertEqual(worksheet["F4"].value, "Farq (Profiles)")
        self.assertEqual(worksheet["F5"].value, 6)

        # Final column is Jami foiz
        self.assertEqual(worksheet.cell(row=3, column=worksheet.max_column).value, "Jami foiz")
        self.assertEqual(worksheet.cell(row=5, column=worksheet.max_column).value, 80)
        self.assertEqual(worksheet.cell(row=6, column=worksheet.max_column).value, 80)
        self.assertNotEqual(
            worksheet.cell(row=5, column=1).fill.fgColor.rgb,
            "00000000",
        )

    def test_exports_with_extra_columns_when_explicitly_requested(self):
        response = self.client.post(
            reverse("mega_excel_export", kwargs={"project_key": "uzchess"}),
            {
                "snapshot_ids": [str(self.first_snapshot.pk), str(self.last_snapshot.pk)],
                "extra_columns": ["certificates_count", "users_ratio_percent"],
                "compare_latest": "on",
                "color_mode": "traffic",
            },
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        worksheet = workbook.active

        # Column B is 14-30 yoshlar
        self.assertEqual(worksheet["B3"].value, "14-30 yoshlar")

        cert_columns = [
            cell.column
            for cell in worksheet[4]
            if cell.value == "Olingan sertifikatlar"
        ]
        self.assertEqual(len(cert_columns), 2)
        self.assertEqual(worksheet.cell(row=5, column=cert_columns[0]).value, 1)
        self.assertEqual(worksheet.cell(row=5, column=cert_columns[1]).value, 3)

        ratio_columns = [
            cell.column
            for cell in worksheet[4]
            if cell.value == "Ishtirokchi / yoshlar (%)"
        ]
        self.assertEqual(len(ratio_columns), 2)
        self.assertEqual(worksheet.cell(row=5, column=ratio_columns[0]).value, 20)
        self.assertEqual(worksheet.cell(row=5, column=ratio_columns[1]).value, 80)

    def test_rejects_invalid_custom_color_thresholds(self):
        response = self.client.post(
            reverse("mega_excel_export", kwargs={"project_key": "uzchess"}),
            {
                "snapshot_ids": [str(self.last_snapshot.pk)],
                "color_mode": "custom",
                "red_limit": "80",
                "yellow_limit": "50",
            },
        )

        self.assertEqual(response.status_code, 400)
        self.assertContains(response, "Qizil chegara sariq chegaradan kichik", status_code=400)
